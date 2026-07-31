#!/usr/bin/env python3
"""export_notes — pull the speaker notes out of a .pptx into a plain-text rehearsal script.

A presenter often rehearses from the notes alone, away from the slides. This reads each
slide's notes (deckkit.speaker_notes writes them) and emits a clean, numbered text file.

Usage:
    python export_notes.py deck.pptx [notes.txt]
    # default out: <deck>.notes.txt  beside the deck
"""
import sys
import os
from pptx import Presentation


def export_notes(pptx_path, out_path=None):
    prs = Presentation(pptx_path)
    if out_path is None:
        out_path = os.path.splitext(pptx_path)[0] + ".notes.txt"
    lines = []
    for i, slide in enumerate(prs.slides, start=1):
        # slide title, if any, for orientation
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        header = f"--- Slide {i}" + (f": {title}" if title else "") + " ---"
        lines.append(header)
        lines.append(notes if notes else "(no notes)")
        lines.append("")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return out_path, len(prs.slides._sldIdLst)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("usage: python export_notes.py deck.pptx [notes.txt]")
        raise SystemExit(2)
    out, n = export_notes(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
    print(f"exported notes for {n} slides -> {out}")
